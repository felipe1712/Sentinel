import openpyxl
import pandas as pd
import numpy as np
import json
from pathlib import Path

excel_file = Path(r"c:\Users\DELL\Documents\SentinelIQ\Sentinelv2\data\Queretaro\electoral\Elecciones Queretaro 15- 18.xlsx")
wb = openpyxl.load_workbook(str(excel_file), data_only=True)

print("=" * 70)
print("1. ANALISIS DETALLADO: GUBERNATURA 2015")
print("=" * 70)

ws2 = wb['Gubernatura 2015']
header2 = [ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)]

rows2 = []
current_dl = 'DESCONOCIDO'
current_mun = 'DESCONOCIDO'

for r in range(2, ws2.max_row + 1):
    c1 = ws2.cell(r, 1).value
    if c1 is None:
        continue
    c1_str = str(c1).strip()
    if 'DISTRITO' in c1_str.upper():
        current_dl = c1_str.split('–')[0].split('-')[0].replace('DISTRITO', '').strip()
        current_mun = c1_str.split('–')[-1].split('-')[-1].strip()
    else:
        try:
            sec = int(float(c1_str))
            row_data = [ws2.cell(r, c).value for c in range(2, ws2.max_column + 1)]
            rows2.append([sec, current_dl, current_mun] + row_data)
        except ValueError:
            pass

df_gub = pd.DataFrame(rows2, columns=['SECCION', 'DISTRITO_LOCAL', 'MUNICIPIO'] + header2[1:])
vote_cols2 = header2[1:]
for c in vote_cols2:
    df_gub[c] = pd.to_numeric(df_gub[c].astype(str).str.replace('-', '0').str.replace(',', '').str.strip(), errors='coerce').fillna(0)

df_gub['VOTOS_PAN'] = df_gub['PAN']

pri_cols = ['PRI', 'PVEM', 'Nueva alianza (NA)', 'PT',
            'PRI-NA-PVEM-PT', 'PRI-NA-PVEM', 'PRI-PVEM-PT', 'PRI-NA-PT', 'NA-PVEM-PT',
            'PRI-NA', 'PRI-PVEM', 'PRI-PT', 'NA-PVEM', 'NA-PT', 'PVEM-PT']
df_gub['VOTOS_PRI_COALICION'] = df_gub[pri_cols].sum(axis=1)
df_gub['VOTOS_MORENA'] = df_gub['Morena']
df_gub['VOTOS_PRD'] = df_gub['PRD']
df_gub['VOTOS_MC'] = df_gub['Movimiento Ciudadano']
df_gub['TOTAL_VOTOS'] = df_gub[vote_cols2].sum(axis=1)

def get_winner_2015(row):
    cand_votes = {
        'PAN': row['VOTOS_PAN'],
        'PRI-PVEM-NA-PT': row['VOTOS_PRI_COALICION'],
        'MORENA': row['VOTOS_MORENA'],
        'PRD': row['VOTOS_PRD'],
        'MC': row['VOTOS_MC']
    }
    sorted_cand = sorted(cand_votes.items(), key=lambda x: x[1], reverse=True)
    w_party, w_votes = sorted_cand[0]
    s_party, s_votes = sorted_cand[1]
    tot = row['TOTAL_VOTOS']
    margin = (w_votes - s_votes) / tot * 100 if tot > 0 else 0
    return pd.Series([w_party, w_votes, s_party, s_votes, margin], index=['GANADOR', 'VOTOS_GANADOR', 'SEGUNDO', 'VOTOS_SEGUNDO', 'MARGEN'])

df_gub[['GANADOR', 'VOTOS_GANADOR', 'SEGUNDO', 'VOTOS_SEGUNDO', 'MARGEN']] = df_gub.apply(get_winner_2015, axis=1)

tot_state = df_gub['TOTAL_VOTOS'].sum()
pan_tot = df_gub['VOTOS_PAN'].sum()
pri_tot = df_gub['VOTOS_PRI_COALICION'].sum()
mor_tot = df_gub['VOTOS_MORENA'].sum()
prd_tot = df_gub['VOTOS_PRD'].sum()
mc_tot = df_gub['VOTOS_MC'].sum()
nul_tot = df_gub['NULOS'].sum()

print(f"Total Secciones computadas: {len(df_gub)}")
print(f"Votos Totales Estatales: {tot_state:,.0f}")
print(f"  PAN (Francisco Domínguez): {pan_tot:,.0f} ({pan_tot/tot_state*100:.2f}%)")
print(f"  PRI-PVEM-NA-PT (Roberto Loyola): {pri_tot:,.0f} ({pri_tot/tot_state*100:.2f}%)")
print(f"  MORENA (Celia Maya): {mor_tot:,.0f} ({mor_tot/tot_state*100:.2f}%)")
print(f"  PRD (Adolfo Camacho): {prd_tot:,.0f} ({prd_tot/tot_state*100:.2f}%)")
print(f"  MC (Salvador López): {mc_tot:,.0f} ({mc_tot/tot_state*100:.2f}%)")
print(f"  Votos Nulos: {nul_tot:,.0f} ({nul_tot/tot_state*100:.2f}%)")

print("\nSecciones ganadas en 2015:")
for party, count in df_gub['GANADOR'].value_counts().items():
    print(f"  {party}: {count} secciones ({count/len(df_gub)*100:.1f}%)")

print("\nResultados por Municipio (Gubernatura 2015):")
mun_summary = df_gub.groupby('MUNICIPIO').agg({
    'SECCION': 'count',
    'VOTOS_PAN': 'sum',
    'VOTOS_PRI_COALICION': 'sum',
    'VOTOS_MORENA': 'sum',
    'TOTAL_VOTOS': 'sum'
}).reset_index()
mun_summary['GANADOR'] = np.where(mun_summary['VOTOS_PAN'] > mun_summary['VOTOS_PRI_COALICION'], 'PAN', 'PRI-PVEM-NA-PT')
mun_summary['PCT_PAN'] = (mun_summary['VOTOS_PAN'] / mun_summary['TOTAL_VOTOS'] * 100).round(2)
mun_summary['PCT_PRI'] = (mun_summary['VOTOS_PRI_COALICION'] / mun_summary['TOTAL_VOTOS'] * 100).round(2)

for _, r in mun_summary.sort_values(by='TOTAL_VOTOS', ascending=False).iterrows():
    print(f"  {r['MUNICIPIO']:<22} | Ganador: {r['GANADOR']:<15} | PAN: {r['VOTOS_PAN']:>7,d} ({r['PCT_PAN']:>5.1f}%) | PRI-Col: {r['VOTOS_PRI_COALICION']:>7,d} ({r['PCT_PRI']:>5.1f}%) | Total: {r['TOTAL_VOTOS']:>8,d} | Secs: {r['SECCION']}")

print("\n" + "=" * 70)
print("2. ANALISIS DETALLADO: DIPUTACIONES 2018")
print("=" * 70)

df_dip_raw = pd.read_excel(str(excel_file), sheet_name='Diputados 2018', skiprows=5)
print(f"Total casillas computadas: {len(df_dip_raw)}")

# Clean numeric columns in 2018
for col in df_dip_raw.columns:
    if col not in ['CLAVE_CASILLA', 'CLAVE_ACTA', 'ID_ESTADO', 'NOMBRE_ESTADO', 'NOMBRE_DISTRITO', 'ID_CASILLA', 'TIPO_CASILLA', 'CASILLA', 'OBSERVACIONES', 'MECANISMOS_TRASLADO', 'FECHA_HORA']:
        df_dip_raw[col] = pd.to_numeric(df_dip_raw[col].astype(str).str.replace('-', '0').str.replace(',', '').str.strip(), errors='coerce').fillna(0)

pan_coal_cols_18 = ['PAN', 'PRD', 'MOVIMIENTO CIUDADANO', 'PAN_PRD_MC', 'PAN_PRD', 'PAN_MC', 'PRD_MC']
morena_coal_cols_18 = ['MORENA', 'PT', 'ENCUENTRO SOCIAL', 'PT_MORENA_PES', 'PT_MORENA', 'PT_PES', 'MORENA_PES']
pri_cols_18 = ['PRI', 'PRI_PVEM_NA', 'PRI_PVEM', 'PRI_NA']
pvem_cols_18 = ['PVEM', 'PVEM_NA']
na_cols_18 = ['NUEVA ALIANZA']

df_dip_raw['VOTOS_PAN_PRD_MC'] = df_dip_raw[pan_coal_cols_18].sum(axis=1)
df_dip_raw['VOTOS_MORENA_PT_PES'] = df_dip_raw[morena_coal_cols_18].sum(axis=1)
df_dip_raw['VOTOS_PRI'] = df_dip_raw[pri_cols_18].sum(axis=1)
df_dip_raw['VOTOS_PVEM'] = df_dip_raw[pvem_cols_18].sum(axis=1)
df_dip_raw['VOTOS_NA'] = df_dip_raw[na_cols_18].sum(axis=1)

# Group by Section
df_sec_18 = df_dip_raw.groupby('SECCION').agg({
    'ID_DISTRITO': 'first',
    'NOMBRE_DISTRITO': 'first',
    'LISTA_NOMINAL_CASILLA': 'sum',
    'TOTAL_VOTOS_CALCULADOS': 'sum',
    'PAN': 'sum',
    'PRI': 'sum',
    'PRD': 'sum',
    'PVEM': 'sum',
    'PT': 'sum',
    'MOVIMIENTO CIUDADANO': 'sum',
    'NUEVA ALIANZA': 'sum',
    'MORENA': 'sum',
    'ENCUENTRO SOCIAL': 'sum',
    'VOTOS_PAN_PRD_MC': 'sum',
    'VOTOS_MORENA_PT_PES': 'sum',
    'VOTOS_PRI': 'sum',
    'VOTOS_PVEM': 'sum',
    'VOTOS_NA': 'sum',
    'VN': 'sum',
    'CNR': 'sum',
    'CLAVE_CASILLA': 'count'
}).reset_index().rename(columns={'CLAVE_CASILLA': 'NUM_CASILLAS'})

def get_winner_2018(row):
    cand_votes = {
        'PAN-PRD-MC': row['VOTOS_PAN_PRD_MC'],
        'MORENA-PT-PES': row['VOTOS_MORENA_PT_PES'],
        'PRI': row['VOTOS_PRI'],
        'PVEM': row['VOTOS_PVEM'],
        'NUEVA ALIANZA': row['VOTOS_NA']
    }
    sorted_cand = sorted(cand_votes.items(), key=lambda x: x[1], reverse=True)
    w_party, w_votes = sorted_cand[0]
    s_party, s_votes = sorted_cand[1]
    tot = row['TOTAL_VOTOS_CALCULADOS']
    margin = (w_votes - s_votes) / tot * 100 if tot > 0 else 0
    return pd.Series([w_party, w_votes, s_party, s_votes, margin], index=['GANADOR', 'VOTOS_GANADOR', 'SEGUNDO', 'VOTOS_SEGUNDO', 'MARGEN'])

df_sec_18[['GANADOR', 'VOTOS_GANADOR', 'SEGUNDO', 'VOTOS_SEGUNDO', 'MARGEN']] = df_sec_18.apply(get_winner_2018, axis=1)

tot_18 = df_sec_18['TOTAL_VOTOS_CALCULADOS'].sum()
ln_18 = df_sec_18['LISTA_NOMINAL_CASILLA'].sum()
pan_18 = df_sec_18['VOTOS_PAN_PRD_MC'].sum()
mor_18 = df_sec_18['VOTOS_MORENA_PT_PES'].sum()
pri_18 = df_sec_18['VOTOS_PRI'].sum()
pve_18 = df_sec_18['VOTOS_PVEM'].sum()
na_18 = df_sec_18['VOTOS_NA'].sum()
vn_18 = df_sec_18['VN'].sum()

print(f"Total Secciones computadas: {len(df_sec_18)}")
print(f"Lista Nominal Estatal: {ln_18:,.0f}")
print(f"Total Votos Calculados: {tot_18:,.0f} (Participacion: {tot_18/ln_18*100:.2f}%)")
print(f"  PAN-PRD-MC: {pan_18:,.0f} ({pan_18/tot_18*100:.2f}%)")
print(f"  MORENA-PT-PES: {mor_18:,.0f} ({mor_18/tot_18*100:.2f}%)")
print(f"  PRI: {pri_18:,.0f} ({pri_18/tot_18*100:.2f}%)")
print(f"  PVEM: {pve_18:,.0f} ({pve_18/tot_18*100:.2f}%)")
print(f"  NUEVA ALIANZA: {na_18:,.0f} ({na_18/tot_18*100:.2f}%)")
print(f"  Votos Nulos: {vn_18:,.0f} ({vn_18/tot_18*100:.2f}%)")

print("\nSecciones ganadas en 2018 (Diputaciones Federales):")
for party, count in df_sec_18['GANADOR'].value_counts().items():
    print(f"  {party}: {count} secciones ({count/len(df_sec_18)*100:.1f}%)")

print("\nResultados por Distrito Federal (Diputaciones 2018):")
dist_summary = df_sec_18.groupby(['ID_DISTRITO', 'NOMBRE_DISTRITO']).agg({
    'SECCION': 'count',
    'NUM_CASILLAS': 'sum',
    'VOTOS_PAN_PRD_MC': 'sum',
    'VOTOS_MORENA_PT_PES': 'sum',
    'VOTOS_PRI': 'sum',
    'TOTAL_VOTOS_CALCULADOS': 'sum',
    'LISTA_NOMINAL_CASILLA': 'sum'
}).reset_index()

dist_summary['GANADOR'] = np.where(dist_summary['VOTOS_PAN_PRD_MC'] > dist_summary['VOTOS_MORENA_PT_PES'], 'PAN-PRD-MC', 'MORENA-PT-PES')
dist_summary['PARTICIPACION'] = (dist_summary['TOTAL_VOTOS_CALCULADOS'] / dist_summary['LISTA_NOMINAL_CASILLA'] * 100).round(2)

for _, r in dist_summary.iterrows():
    print(f"  Distrito {int(r['ID_DISTRITO'])} ({r['NOMBRE_DISTRITO']}): Ganador = {r['GANADOR']} | PAN-Frente: {r['VOTOS_PAN_PRD_MC']:>7,d} | MORENA-JHH: {r['VOTOS_MORENA_PT_PES']:>7,d} | PRI: {r['VOTOS_PRI']:>6,d} | Part: {r['PARTICIPACION']}% | Secs: {r['SECCION']}")
